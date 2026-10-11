# inventory_gaps.py — 量化"说明书跟着变"这件事还差什么
import collections, csv, html, importlib.util, json, os, re
import openpyxl

APPLY = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location("ug", os.path.join(APPLY, "update_guide.py"))
ug = importlib.util.module_from_spec(spec)
spec.loader.exec_module(ug)

MOD = os.path.join(os.path.dirname(ug.ROOT), "法魂mod V3.0.3 汉化本体（适配游戏版本1.17.1）（新YAFSML引擎V0.10.4）")
XLSX = os.path.join(ug.ROOT, "法魂资料整合表.xlsx")


def main():
    src = open(os.path.join(APPLY, "work", "msg", "zh", "GoodsCaption.fmg.xml"), encoding="utf-8-sig").read()

    wb0 = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    setnames, cur0 = set(), ""
    for r in wb0["套装表"].iter_rows(values_only=True):
        v = ["" if x is None else str(x).strip() for x in r]
        if not any(v):
            continue
        if v[0]:
            cur0 = v[0]
            setnames.add(cur0)
    lines_all, set_lines, item_lines = 0, [], []
    for eid in ug.GUIDE_ENTRIES:
        m = re.search(r'<text id="%d">(.*?)</text>' % eid, src, re.S)
        if not m:
            continue
        for _x in ug.LINE.finditer(html.unescape(m.group(1))):
            a, b = _x.group(2), _x.group(5)
            lines_all += 1
            nm = a.strip()
            is_set = nm.endswith("套装") or nm in setnames
            (set_lines if is_set else item_lines).append((eid, nm, b.strip()))
    print("=== ① 说明书行构成（条目 8800-8826 / 8870-8895）===")
    print(f"  总行数 {lines_all} | 物品行 {len(item_lines)} | 套装行 {len(set_lines)}")
    print(f"  套装行样例: {[x[1] for x in set_lines[:8]]}")

    full = [r for r in csv.DictReader(open(os.path.join(APPLY, "out", "plan_guide_default.csv"), encoding="utf-8-sig"))
            if r["old"] != r["new"]]
    names = {"protector": ug.fmg(os.path.join(ug.CAT, "evidence", "xml", "ProtectorName.fmg.xml"),
                                 os.path.join(ug.CAT, "evidence", "xml", "ProtectorName_dlc01.fmg.xml")),
             "weapon": ug.fmg(os.path.join(ug.WB, "evidence", "xml", "WeaponName.fmg.xml"),
                              os.path.join(ug.WB, "evidence", "xml", "WeaponName_dlc01.fmg.xml")),
             "goods": ug.fmg(os.path.join(ug.CAT, "evidence", "xml", "GoodsName.fmg.xml"),
                             os.path.join(ug.CAT, "evidence", "xml", "GoodsName_dlc01.fmg.xml"))}
    name2id = {}
    for c, t in names.items():
        for i, n in t.items():
            name2id.setdefault(n, f"{c}:{i}")
    involved = {r["old"] for r in full} | {r["new"] for r in full}
    hit_item = [(e, a, b) for e, a, b in item_lines if name2id.get(a) in involved]
    print("\n=== ② 完整包（750 行 / 64 池）会牵动的说明书行 ===")
    print(f"  涉及的物品 {len(involved)} 个 | 说明书里的物品行 {len(item_lines)} 行，其中受影响 {len(hit_item)} 行（逻辑已支持）")

    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    setmap, cur = {}, ""
    for r in wb["套装表"].iter_rows(values_only=True):
        v = ["" if x is None else str(x).strip() for x in r]
        if not any(v):
            continue
        if v[0]:
            cur = v[0]
        if len(v) > 1 and v[1]:
            base = cur.replace("套装", "").replace("套", "")
            for k in {cur, base + "套装", base + "套"}:
                setmap.setdefault(k, []).append(v[1])
    affected = []
    for eid, a, b in set_lines:
        ids = {name2id.get(p) for p in setmap.get(a, []) if name2id.get(p)}
        if ids & involved:
            affected.append((eid, a, b, len(ids)))
    print(f"\n=== ③ 套装行（目前**不会改**）===")
    print(f"  被完整包牵动的套装行 {len(affected)} 行 / 共 {len(set_lines)} 行")
    for eid, a, b, n in affected[:12]:
        print(f"     {eid}  【{a}】 - {b}   （套装表里 {n} 件，其中若干在池里）")

    msgdir = os.path.join(MOD, "mod", "msg")
    langs = [d for d in sorted(os.listdir(msgdir)) if os.path.exists(os.path.join(msgdir, d, "item_dlc02.msgbnd.dcx"))]
    print(f"\n=== ④ 语言版本 ===")
    print(f"  带 item_dlc02.msgbnd.dcx 的语言 {len(langs)} 个: {langs}")
    print("  已更新: 只有 zhocn")

    rep = json.load(open(os.path.join(APPLY, "logs", "guide_apply_pilot.json"), encoding="utf-8"))
    pm = rep.get("place_miss", [])
    cats = collections.Counter()
    for old, new, tbl, rid, rn in pm:
        cats["商店行" if tbl == "ShopLineupParam" else ("行名没有方括号" if "[" not in (rn or "") else "括号里没有可匹配地名")] += 1
    print(f"\n=== ⑤ 认不出地点的行（首测 {len(pm)} 行）===")
    print(f"  分类: {dict(cats)}")
    for old, new, tbl, rid, rn in pm:
        print(f"     {old:<20} → {new:<20} | {tbl}:{rid:<10} | {rn!r}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
