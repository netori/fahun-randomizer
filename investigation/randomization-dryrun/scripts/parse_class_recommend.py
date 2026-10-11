# parse_class_recommend.py — 只读：解析《职业推荐道具表3.0》与《套装表》，给护甲建立"职业归属"
# 目的：修正"未归属装备混入池子"的错误（例：大山羊臂甲）
#   职业推荐道具表3.0：职业 | 类别(装备/武器/卢恩) | 名称 | 获取位置  → 职业↔物品 归属
#   套装表：套装名称 | 装备名称 | 属性效果 | 获取地点            → 权威的"成套"识别
# 输出：investigation/randomization-dryrun/evidence/class_attribution.json
import collections, json, os, sys
import openpyxl

DRY = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WS = os.path.dirname(os.path.dirname(os.path.dirname(DRY.split(os.sep + "investigation")[0])))  # 工作区
WS = os.path.join(os.path.dirname(os.path.dirname(DRY)), "")  # 随机mod 目录
WS = os.path.dirname(os.path.dirname(DRY))                    # 随机mod


def main():
    xlsx = os.path.join(WS, "法魂资料整合表.xlsx")
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)

    # ---- 1) 职业推荐道具表3.0 ----
    ws = wb["职业推荐道具表3.0"]
    rows = [["" if c is None else str(c).strip() for c in r] for r in ws.iter_rows(values_only=True)]
    rec = collections.defaultdict(lambda: collections.defaultdict(list))   # class -> 类别 -> [名称]
    cur_cls, cur_kind = "", ""
    for r in rows[1:]:
        if not any(r):
            continue
        if r[0]:
            cur_cls = r[0]
        if len(r) > 1 and r[1]:
            cur_kind = r[1]
        name = r[2] if len(r) > 2 else ""
        loc = r[3] if len(r) > 3 else ""
        if name:
            rec[cur_cls][cur_kind or "(未标类别)"].append({"name": name, "where": loc})
    print("=== 职业推荐道具表3.0 ===")
    kinds = collections.Counter(k for c in rec.values() for k in c)
    print("  职业数:", len(rec), "| 类别分布:", dict(kinds))
    tot = sum(len(v) for c in rec.values() for v in c.values())
    print("  条目总数:", tot)
    for cls in list(rec)[:3]:
        print(f"    {cls}: " + "; ".join(f"{k}×{len(v)}" for k, v in rec[cls].items()))

    # 物品名 → 职业（反向索引）
    name2cls = collections.defaultdict(set)
    for cls, bykind in rec.items():
        for kind, items in bykind.items():
            for it in items:
                name2cls[it["name"]].add(cls)

    # ---- 2) 套装表 ----
    ws2 = wb["套装表"]
    rows2 = [["" if c is None else str(c).strip() for c in r] for r in ws2.iter_rows(values_only=True)]
    sets = collections.defaultdict(list)
    cur_set = ""
    for r in rows2[1:]:
        if not any(r):
            continue
        if r[0]:
            cur_set = r[0]
        if len(r) > 1 and r[1]:
            sets[cur_set].append({"name": r[1], "effect": r[2] if len(r) > 2 else "", "where": r[3] if len(r) > 3 else ""})
    print("\n=== 套装表 ===")
    print("  套装数:", len(sets), "| 四件以上的套装:", sum(1 for v in sets.values() if len(v) >= 4))
    for s in list(sets)[:3]:
        print(f"    {s}: {[x['name'] for x in sets[s]]}")
    name2set = collections.defaultdict(set)
    for s, items in sets.items():
        for it in items:
            name2set[it["name"]].add((s, len(items)))

    # ---- 3) 检查大山羊 与 商店护甲清单的归属覆盖 ----
    print("\n=== 关键排查：大山羊臂甲 ===")
    for probe in ("大山羊臂甲", "大山羊头盔", "大山羊铠甲", "大山羊腿甲"):
        print(f"  {probe}: 职业推荐表命中={sorted(name2cls.get(probe, [])) or '无'} | 套装表命中={sorted(name2set.get(probe, [])) or '无'}")
    import csv
    draft = list(csv.DictReader(open(os.path.join(os.path.dirname(DRY), "randomization-catalog-review", "shop_candidate_draft.csv"),
                                     encoding="utf-8-sig")))
    armor = [r for r in draft if r["category"] == "protector"]
    have_cls = [r for r in armor if name2cls.get(r["name_cn"])]
    have_set = [r for r in armor if name2set.get(r["name_cn"])]
    print(f"\n=== 商店护甲清单（{len(armor)} 件）的归属覆盖 ===")
    print(f"  职业推荐表有归属: {len(have_cls)} 件；套装表有归属: {len(have_set)} 件；两者都没有: "
          f"{sum(1 for r in armor if not name2cls.get(r['name_cn']) and not name2set.get(r['name_cn']))} 件")
    print("  有职业归属的（前 12）:", [(r["item_id"], r["name_cn"], sorted(name2cls[r["name_cn"]])[:2]) for r in have_cls[:12]])
    print("  无归属的（前 12）:", [(r["item_id"], r["name_cn"]) for r in armor
                                 if not name2cls.get(r["name_cn"]) and not name2set.get(r["name_cn"])][:12])

    out = {"name_to_classes": {k: sorted(v) for k, v in name2cls.items()},
           "name_to_set": {k: sorted(v) for k, v in name2set.items()},
           "sets": {s: [x["name"] for x in v] for s, v in sets.items()},
           "class_recommend_counts": {c: {k: len(v) for k, v in bk.items()} for c, bk in rec.items()}}
    json.dump(out, open(os.path.join(DRY, "evidence", "class_attribution.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("\n写出 evidence/class_attribution.json")
    return 0


if __name__ == "__main__":
    sys.exit(main())
