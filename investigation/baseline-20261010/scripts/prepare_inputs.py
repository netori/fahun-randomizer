"""Copy only three primary archives; inspect only worksheet 0 using bundled openpyxl."""
from common import *
import openpyxl, csv, collections, re

sources = [
    (MOD / 'mod/regulation.bin', BASE / 'inputs/regulation.bin', 'active-name regulation candidate; original not tool-processed'),
    (MOD / 'mod/msg/zhocn/item_dlc02.msgbnd.dcx', BASE / 'inputs/zhocn/item_dlc02.msgbnd.dcx', 'simplified Chinese item text container'),
    (MOD / 'mod/msg/zhocn/menu_dlc02.msgbnd.dcx', BASE / 'inputs/zhocn/menu_dlc02.msgbnd.dcx', 'simplified Chinese menu text container'),
]
copies = []
for source, target, role in sources:
    copies.append(copy_registered(source, target, role))
    with source.open('rb') as stream:
        header = stream.read(128)
    print(json.dumps({'source': str(source), 'size': source.stat().st_size, 'sha256': copies[-1]['sha256'], 'first_128_bytes_hex': header.hex()}, ensure_ascii=False))
save_json(BASE / 'evidence/input_copies.json', copies)
# Version text is documentary evidence only, never substitute for internal PARAM version.
for p in [MOD / 'version.txt', MOD / 'Start_Convergence_YAFSML.bat']:
    if p.is_file():
        print('registered documentary source:', json.dumps(record_source(p, 'MOD documentary version/launcher path evidence'), ensure_ascii=False))

book = WORKSPACE / '法魂资料整合表.xlsx'
record_source(book, 'openpyxl worksheet zero inspection')
wb = openpyxl.load_workbook(book, read_only=True, data_only=False, keep_links=False)
ws = wb.worksheets[0]
rows = []
for num, values in enumerate(ws.iter_rows(), 1):
    populated = [{'column': c.column, 'coordinate': c.coordinate, 'type': c.data_type, 'value': c.value} for c in values if c.value is not None]
    if populated:
        rows.append({'excel_row': num, 'cells': populated})
header = {cell['column']: str(cell['value']) for cell in rows[0]['cells']}
entries = []
for row in rows[1:]:
    vals = {c['column']: c['value'] for c in row['cells']}
    if any(v is not None and str(v).strip() for v in vals.values()):
        entries.append({'excel_row': row['excel_row'], **{header.get(col, f'column_{col}'): vals.get(col) for col in sorted(header)}})
formulas = [c['coordinate'] for r in rows for c in r['cells'] if c['type'] == 'f']
numeric_cells = [c for r in rows[1:] for c in r['cells'] if isinstance(c['value'], (int, float)) and not isinstance(c['value'], bool)]
summary = {'workbook_absolute_path': str(book.resolve()), 'worksheet_index': 0, 'worksheet_name': ws.title, 'sheet_names_metadata': wb.sheetnames, 'max_row': ws.max_row, 'max_column': ws.max_column, 'header': header, 'nonempty_data_rows': len(entries), 'empty_location_rows': sum(1 for e in entries if not str(e.get('Location') or '').strip()), 'formulas_first_sheet': formulas, 'numeric_cell_count_first_sheet': len(numeric_cells), 'merged_ranges': 'not materialized in read_only mode; source unchanged', 'openpyxl_version': openpyxl.__version__, 'other_sheets_values_read': False}
save_json(BASE / 'evidence/excel_first_sheet_summary.json', summary)
save_json(BASE / 'evidence/excel_first_sheet_rows.json', entries)
with output_path(BASE / 'evidence/excel_first_sheet.csv').open('x', encoding='utf-8-sig', newline='') as stream:
    writer = csv.DictWriter(stream, fieldnames=['excel_row'] + [header[c] for c in sorted(header)])
    writer.writeheader()
    writer.writerows(entries)
print('excel_summary:', json.dumps(summary, ensure_ascii=False))
pattern = re.compile('套装|说明|指南|卢恩|幽影|商人|商店|三件|改装|位置')
samples = [e for e in entries if pattern.search(str(e.get('Name') or ''))][:30]
print('name samples only:', json.dumps(samples, ensure_ascii=False))
wb.close()
